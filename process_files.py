"""
Day 3: Excel, Word and Python File Processing
Run:
    pip install -r requirements.txt
    python process_files.py

This script reads the sample club-registration Excel workbook, cleans the
data, writes an organized Excel workbook, and creates a Word report.
"""
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "club_registration_data.xlsx"
OUTPUT_EXCEL = BASE_DIR / "organized_club_registrations.xlsx"
OUTPUT_WORD = BASE_DIR / "club_registration_cleaning_report.docx"

DEPARTMENT_NAMES = {
    "COMPUTER SCIENCE": "CSE",
    "CSE": "CSE",
    "INFORMATION TECHNOLOGY": "IT",
    "IT": "IT",
    "DATA SCIENCE": "DS",
    "DS": "DS",
}

def clean_data():
    """Read Excel, normalize text, remove duplicates, and organize rows."""
    data = pd.read_excel(INPUT_FILE, sheet_name=0)
    original_count = len(data)

    # Standardize column labels and remove unwanted spaces from text fields.
    data.columns = [str(column).strip().replace(" ", "_") for column in data.columns]
    for column in data.select_dtypes(include=["object", "string"]).columns:
        data[column] = data[column].astype("string").str.strip()

    # Standardize email values and department names.
    data["Email"] = data["Email"].astype("string").str.strip().str.lower()
    data["Email"] = data["Email"].replace("", pd.NA)
    data["Department"] = (
        data["Department"].astype("string").str.strip().str.upper()
        .map(lambda value: DEPARTMENT_NAMES.get(value, value) if pd.notna(value) else value)
    )

    # Keep the first occurrence of each non-empty email address.
    has_email = data["Email"].notna()
    duplicate_mask = has_email & data.duplicated(subset=["Email"], keep="first")
    duplicate_count = int(duplicate_mask.sum())
    data = data.loc[~duplicate_mask].copy()

    # Add a clear review status and sort the final records.
    data["Review_Status"] = data["Email"].apply(
        lambda value: "Needs email review" if pd.isna(value) or not str(value).strip() else "OK"
    )
    data = data.sort_values(["Department", "Student_Name"], na_position="last").reset_index(drop=True)

    # Save the cleaned dataset as a formatted Excel workbook.
    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name="Organized Registrations")
        sheet = writer.sheets["Organized Registrations"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 28
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 3, 30)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    return data, original_count, duplicate_count

def create_word_report(data, original_count, duplicate_count):
    """Create a Word report directly from the cleaned data."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_heading("Club Registration Data Cleaning Report", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(15, 118, 110)
        run.font.size = Pt(22)

    doc.add_paragraph(
        "This report was created automatically by Python. The program reads the sample "
        "Excel registration file, standardizes the entries, saves an organized workbook, "
        "and generates this Word report without manually opening either application."
    )

    doc.add_heading("Processing summary", level=1)
    summary = doc.add_table(rows=0, cols=2)
    summary.style = "Light Shading Accent 1"
    items = [
        ("Rows in source file", str(original_count)),
        ("Duplicate registrations removed", str(duplicate_count)),
        ("Rows in organized workbook", str(len(data))),
        ("Rows needing email review", str((data["Review_Status"] != "OK").sum())),
        ("Excel output", OUTPUT_EXCEL.name),
    ]
    for key, value in items:
        cells = summary.add_row().cells
        cells[0].text = key
        cells[1].text = value

    doc.add_heading("Cleaning steps completed", level=1)
    steps = [
        "Removed leading and trailing spaces from text values.",
        "Converted email addresses to lowercase and identified missing emails.",
        "Standardized department names to CSE, IT, or DS.",
        "Removed repeated registrations using email address while keeping the first entry.",
        "Sorted records by department and student name.",
        "Saved the updated data to a separate Excel workbook.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("Organized registration preview", level=1)
    preview_cols = ["Club_ID", "Student_Name", "Department", "Year", "Review_Status"]
    table = doc.add_table(rows=1, cols=len(preview_cols))
    table.style = "Table Grid"
    for i, col in enumerate(preview_cols):
        table.rows[0].cells[i].text = col.replace("_", " ")
    for _, row in data.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(preview_cols):
            value = row[col]
            cells[i].text = "" if pd.isna(value) else str(value)

    doc.add_paragraph(
        "Sample data note: all names, email addresses, and phone numbers in this project are fictional."
    )
    doc.save(OUTPUT_WORD)

if __name__ == "__main__":
    cleaned, source_count, duplicates = clean_data()
    create_word_report(cleaned, source_count, duplicates)
    print("Day 3 file processing completed successfully!")
    print(f"Source registrations: {source_count}")
    print(f"Duplicate registrations removed: {duplicates}")
    print(f"Organized registrations saved: {len(cleaned)}")
    print(f"Excel output: {OUTPUT_EXCEL.name}")
    print(f"Word output: {OUTPUT_WORD.name}")
